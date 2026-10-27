from pathlib import Path
import io
import fitz
from docx import Document

def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = content.decode("utf-8-sig", errors="replace")
    elif suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            text = "\\n".join(page.get_text("text") for page in pdf)
    elif suffix == ".docx":
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = [" | ".join(c.text.strip() for c in row.cells) for table in doc.tables for row in table.rows]
        text = "\\n".join(paragraphs + tables)
    else:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")
    if not text.strip():
        raise ValueError("No text extracted. File may be scanned; OCR is not enabled.")
    return text
